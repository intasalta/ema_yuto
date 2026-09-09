import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Lista explícita de archivos a convertir
ARCHIVOS_A_CONVERTIR = ["downld02.txt", "downld08.txt"]

COL_NAMES = [
    "Fecha", "Hora", "Temp Ext (°C)", "Temp Máx", "Temp Mín", "Humedad Ext (%)", "Punto Rocío",
    "Vel Vent", "Dir Vent", "Ráfaga Vent", "Vel Máx", "Dir Máx", "Sens Term Wind",
    "Índice Calor", "THW", "THSW", "Presión (hPa)", "Lluvia (mm)", "Int Lluvia",
    "Rad Solar", "Energía Solar", "Rad Solar Máx", "UV", "Dosis UV", "UV Máx",
    "Grados Día C", "Grados Día F", "Temp Int", "Humedad Int", "ET", "Muestras Vent",
    "Tx Vent", "Recepción ISS", "Intervalo Arc"
]

header_fill = PatternFill(start_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
data_font = Font(name="Calibri", size=10)
zebra_fill = PatternFill(start_color="F2F7FA", fill_type="solid")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
)

for file_path in ARCHIVOS_A_CONVERTIR:
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            lines = f.readlines()
    except FileNotFoundError:
        print(f"Archivo no encontrado: {file_path}")
        continue

    data_rows = []
    for line in lines[3:]:
        parts = line.strip().split()
        if len(parts) >= 30:
            data_rows.append(parts)

    df = pd.DataFrame(data_rows)
    if len(df.columns) == len(COL_NAMES):
        df.columns = COL_NAMES

    base_name = file_path.rsplit('.', 1)[0]
    
    # Exportar CSV
    df.to_csv(f"{base_name}.csv", index=False)

    # Exportar Excel formateado
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos Meteorológicos"

    ws.merge_cells("A1:AH1")
    ws["A1"] = f"Reporte Estación Meteorológica - {base_name}"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E78")

    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=3, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, row in df.iterrows():
        row_num = r_idx + 4
        fill = zebra_fill if r_idx % 2 == 1 else PatternFill(fill_type=None)
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=c_idx)
            try:
                cell.value = float(val)
                cell.number_format = "0.0" if "." in str(val) else "0"
            except (ValueError, TypeError):
                cell.value = str(val) if val else ""
            
            cell.font = data_font
            if fill.fill_type: cell.fill = fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 13

    ws.freeze_panes = "A4"
    wb.save(f"{base_name}.xlsx")
    print(f"Procesado: {file_path} -> {base_name}.csv y {base_name}.xlsx")